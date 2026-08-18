"""Scrittura Excel formattata con openpyxl.

Aggiunge ai fogli generati:
- prima riga (header) BLOCCATA quando si scorre (freeze panes A2);
- header in grassetto su sfondo colorato;
- larghezze colonne in base al contenuto (con tetto massimo);
- filtro automatico su tutte le colonne;

opzioni extra:
- km_cols        : colonne mostrate con unità km  (es. '0.00 "km"');
- decimal_cols   : colonne numeriche con molti decimali (es. lat/long);
- red_cols       : in queste colonne, le celle con valore == red_value
                   vengono evidenziate in rosso;
- red_value      : valore sentinella (default -100);
- area_col       : colonna dell'area -> righe della stessa area hanno lo
                   stesso colore di sfondo, aree diverse in alternanza.
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

RED_FILL = PatternFill("solid", fgColor="FFC7CE")     # rosso chiaro
RED_FONT = Font(color="9C0006", bold=True)            # testo rosso scuro
AREA_FILLS = [
    PatternFill("solid", fgColor="E7F0FB"),           # azzurro chiaro
    PatternFill("solid", fgColor="FFFFFF"),           # bianco
]

MAX_COL_WIDTH = 45
MIN_COL_WIDTH = 8

KM_FORMAT = '0.00 "km"'
DEC_FORMAT = "0.000000"


def write_formatted_excel(
    df,
    path: str,
    sheet_name: str = "Dati",
    km_cols: tuple = (),
    decimal_cols: tuple = (),
    red_cols: tuple = (),
    red_value: float = -100.0,
    area_col: str | None = None,
) -> None:
    """Scrive il DataFrame in `path` con header bloccato e formattazione."""
    # Prima scrive il file con pandas (gestisce tipi/date), poi lo stila
    df.to_excel(path, index=False, sheet_name=sheet_name)

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name]
    ncols = ws.max_column
    nrows = ws.max_row

    # mappa nome colonna -> indice (dalla riga header)
    header_names = {c.value: i for i, c in enumerate(ws[1], start=1)}

    # 1) Header bloccato: freeze alla riga 2 (riga 1 sempre visibile)
    ws.freeze_panes = "A2"

    # 2) Stile header
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    # 3) Colore di sfondo alternato per area (righe della stessa area insieme)
    if area_col and area_col in header_names:
        acol = header_names[area_col]
        current_area = None
        fill_idx = 0
        for row in range(2, nrows + 1):
            val = ws.cell(row=row, column=acol).value
            if val != current_area:
                current_area = val
                fill_idx = 1 - fill_idx  # alterna quando cambia area
            fill = AREA_FILLS[fill_idx]
            for col in range(1, ncols + 1):
                ws.cell(row=row, column=col).fill = fill

    # 4) Formato numeri: km e decimali
    for name in km_cols:
        if name in header_names:
            c = header_names[name]
            for row in range(2, nrows + 1):
                cell = ws.cell(row=row, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = KM_FORMAT
    for name in decimal_cols:
        if name in header_names:
            c = header_names[name]
            for row in range(2, nrows + 1):
                cell = ws.cell(row=row, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = DEC_FORMAT

    # 5) Evidenzia in rosso le celle col valore sentinella (es. -100)
    for name in red_cols:
        if name in header_names:
            c = header_names[name]
            for row in range(2, nrows + 1):
                cell = ws.cell(row=row, column=c)
                if cell.value == red_value:
                    cell.fill = RED_FILL
                    cell.font = RED_FONT

    # 6) Larghezze colonne basate sul contenuto (tetto MAX_COL_WIDTH)
    for col in range(1, ncols + 1):
        max_len = len(str(ws.cell(row=1, column=col).value or ""))
        for row in range(2, nrows + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(col)].width = max(
            MIN_COL_WIDTH, min(max_len + 2, MAX_COL_WIDTH)
        )

    # 7) Filtro automatico sull'intera tabella
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"

    wb.save(path)
