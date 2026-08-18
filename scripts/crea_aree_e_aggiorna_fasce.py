"""Script per creare la collection 'aree' in Firebase e aggiornare 'fasceOrarie'.

Legge il file Excel export_informatorems_aree.xlsx (foglio 'Dati') e:

1. Crea/popola la collection 'aree':
   - documentId = chiave formattata da codArea (colonna U): 'A' + numero a
     3 cifre (es. 7 -> 'A007'); il sentinel -1 (non geolocalizzato) -> '0000'
   - campo 'Area'  = NomeArea (colonna V), incluso 'NON GEOLOCALIZZATO'

2. Aggiorna i documenti della collection 'fasceOrarie' (identificati da
   'idFascia', colonna A), aggiungendo/aggiornando i campi:
   - 'idArea' : stessa chiave formattata usata come documentId di 'aree'
                (es. 'A007' per codArea 7, '0000' per codArea -1), cosi' il
                collegamento area->fascia e' diretto
   - 'lat'    : colonna R (Latitudine), valore numerico (-100 se non geolocalizzato)
   - 'lng'    : colonna S (Longitudine), valore numerico (-100 se non geolocalizzato)

Mappatura colonne Excel (1-based):
  A(1)  = IdFascia        R(18) = Latitudine
  U(21) = codArea         S(19) = Longitudine
  V(22) = NomeArea

Per un run di prova senza scrivere su Firebase: impostare DRY_RUN=1.
"""
from __future__ import annotations

import os
import sys

import firebase_admin
import openpyxl
from firebase_admin import credentials
from firebase_admin import firestore


SERVICE_ACCOUNT_PATH = os.environ.get(
    "GOOGLE_APPLICATION_CREDENTIALS",
    "scripts/serviceAccountKey.json",
)
PROJECT_ID = "informatorems-784a1"

EXCEL_PATH = os.environ.get(
    "EXPORT_AREAS_EXCEL", "scripts/export_informatorems_aree.xlsx"
)
SHEET_NAME = "Dati"

AREE_COLLECTION = "aree"
FASCIA_COLLECTION = "fasceOrarie"

BATCH_SIZE = 450

# Colonna Excel (1-based) -> nostra chiave
COL_ID_FASCIA = 1
COL_LAT = 18
COL_LNG = 19
COL_ID_AREA = 21
COL_NOME_AREA = 22


def format_area_id(cod_area) -> str | None:
    """Ricompre il codArea Excel nella chiave documento di 'aree'.

    - codArea -1 (non geolocalizzato) -> '0000'
    - altrimenti -> 'A' + numero a 3 cifre (es. 7 -> 'A007', 12 -> 'A012')

    La stessa chiave viene usata come documentId di 'aree' e come campo
    'idArea' delle 'fasceOrarie', per avere un collegamento diretto.
    """
    if cod_area is None or str(cod_area).strip() == "":
        return None
    n = int(cod_area)
    if n == -1:
        return "0000"
    return f"A{n:03d}"


def load_rows(path: str) -> list[dict]:
    """Legge il foglio 'Dati' e restituisce una riga per ogni IdFascia presente."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]

    rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        id_fascia = ws.cell(r, COL_ID_FASCIA).value
        if id_fascia is None or not str(id_fascia).strip():
            continue
        id_area = ws.cell(r, COL_ID_AREA).value
        nome_area = ws.cell(r, COL_NOME_AREA).value
        lat = ws.cell(r, COL_LAT).value
        lng = ws.cell(r, COL_LNG).value
        rows.append({
            "idFascia": str(id_fascia).strip(),
            "idArea": format_area_id(id_area),
            "Area": str(nome_area) if nome_area is not None else "",
            "lat": lat,
            "lng": lng,
        })
    return rows


def get_db():
    cred = credentials.Certificate(SERVICE_ACCOUNT_PATH)
    firebase_admin.initialize_app(cred, {"projectId": PROJECT_ID})
    return firestore.client()


def upsert_aree(db, rows: list[dict], dry_run: bool) -> int:
    """Crea/aggiorna un documento 'aree' per ogni idArea distinto.

    documentId = idArea; campo 'Area' = NomeArea (il primo valore incontrato,
    coerente tra le righe). Ritorna il numero di aree scritte.
    """
    # Mantieni il primo NomeArea incontrato per idArea (l'ordine e' stabile
    # e la corrispondenza idArea->Area e' univoca).
    aree_map: dict[str, str] = {}
    for row in rows:
        id_area = row["idArea"]
        if id_area is None:
            continue
        if id_area not in aree_map:
            aree_map[id_area] = row["Area"]

    batch = db.batch()
    counter = 0
    for id_area, nome in sorted(aree_map.items()):
        ref = db.collection(AREE_COLLECTION).document(id_area)
        if dry_run:
            print(f"  [DRY] aree/{id_area}: Area={nome!r}")
        else:
            batch.set(ref, {"Area": nome})
            counter += 1
            if counter >= BATCH_SIZE:
                batch.commit()
                batch = db.batch()
                counter = 0
    if counter > 0 and not dry_run:
        batch.commit()

    return len(aree_map)


def update_fasce(db, rows: list[dict], dry_run: bool) -> tuple[int, int]:
    """Aggiorna le fasceOrarie impostando idArea, lat, lng.

    Ritorna (aggiornate, non_trovate). Usa update() cosi' gli altri campi del
    documento vengono preservati.
    """
    batch = db.batch()
    counter = 0
    aggiornate = 0
    non_trovate = 0

    for row in rows:
        id_fascia = row["idFascia"]
        payload: dict = {}
        if row["idArea"] is not None:
            payload["idArea"] = row["idArea"]
        if row["lat"] is not None:
            payload["lat"] = row["lat"]
        if row["lng"] is not None:
            payload["lng"] = row["lng"]
        if not payload:
            continue

        if dry_run:
            print(f"  [DRY] fasceOrarie/{id_fascia}: {payload}")
            aggiornate += 1
            continue

        doc_ref = db.collection(FASCIA_COLLECTION).document(id_fascia)
        if not doc_ref.get().exists:
            print(f"  [WARNING] Fascia non trovata in Firebase: {id_fascia}")
            non_trovate += 1
            continue

        batch.update(doc_ref, payload)
        counter += 1
        aggiornate += 1
        if counter >= BATCH_SIZE:
            batch.commit()
            batch = db.batch()
            counter = 0

    if counter > 0 and not dry_run:
        batch.commit()

    return aggiornate, non_trovate


def main() -> int:
    dry_run = os.environ.get("DRY_RUN") == "1"

    if not os.path.exists(EXCEL_PATH):
        print(f"ERRORE: file non trovato: {EXCEL_PATH}")
        return 1

    print(f"Lettura file Excel: {EXCEL_PATH}")
    rows = load_rows(EXCEL_PATH)
    if not rows:
        print("ERRORE: nessuna riga con IdFascia trovata.")
        return 1
    print(f"Righe totali (una per IdFascia): {len(rows)}")

    print("\nConnessione a Firebase...")
    try:
        db = get_db()
    except Exception as e:
        print(f"ERRORE durante l'inizializzazione Firebase: {e}")
        return 1

    print(f"\n[1/2] Upsert collection '{AREE_COLLECTION}'...")
    n_aree = upsert_aree(db, rows, dry_run)
    print(f"  Aree scritte (documentId = chiave formattata, es. 'A007'/'0000'): {n_aree}")

    print(f"\n[2/2] Aggiornamento collection '{FASCIA_COLLECTION}'...")
    aggiornate, non_trovate = update_fasce(db, rows, dry_run)
    print(f"  Fasce aggiornate (idArea/lat/lng): {aggiornate}")
    print(f"  Non trovate: {non_trovate}")

    print("\nRisultati:")
    print(f"  Aree in '{AREE_COLLECTION}': {n_aree}")
    print(f"  Fasce aggiornate in '{FASCIA_COLLECTION}': {aggiornate}")
    if dry_run:
        print("\n  (DRY RUN: nessuna scrittura effettuata. Per scrivere, rimuovere DRY_RUN=1)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
